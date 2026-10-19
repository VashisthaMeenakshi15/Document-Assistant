import os
import tempfile
import pandas as pd
from pptx import Presentation
from pypdf import PdfReader
from docx import Document as DocxDocument
from bs4 import BeautifulSoup

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter


def parse_documents(uploaded_files, chunk_size: int = 500, chunk_overlap: int = 100):
    """
    Parses PPTX, PDF, DOCX, HTML, Excel, and CSV files, attaching location metadata,
    and returns chunked LangChain Document objects.
    """
    all_docs = []

    for uploaded_file in uploaded_files:
        filename = uploaded_file.name
        file_extension = os.path.splitext(filename)[1].lower()

        # Save bytes to temporary file for parsers
        with tempfile.NamedTemporaryFile(delete=False, suffix=file_extension) as tmp_file:
            tmp_file.write(uploaded_file.getvalue())
            tmp_path = tmp_file.name

        try:
            # --- 1. PowerPoint (.pptx) ---
            if file_extension == ".pptx":
                prs = Presentation(tmp_path)
                for slide_idx, slide in enumerate(prs.slides, start=1):
                    slide_text = []
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            for paragraph in shape.text_frame.paragraphs:
                                text = paragraph.text.strip()
                                if text:
                                    slide_text.append(text)

                    if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                        notes = slide.notes_slide.notes_text_frame.text.strip()
                        if notes:
                            slide_text.append(f"[Speaker Notes: {notes}]")

                    full_content = "\n".join(slide_text)
                    if full_content.strip():
                        all_docs.append(Document(
                            page_content=full_content,
                            metadata={"source": filename, "location_type": "Slide", "location_number": slide_idx}
                        ))

            # --- 2. PDF (.pdf) ---
            elif file_extension == ".pdf":
                reader = PdfReader(tmp_path)
                for page_idx, page in enumerate(reader.pages, start=1):
                    text = page.extract_text()
                    if text and text.strip():
                        all_docs.append(Document(
                            page_content=text.strip(),
                            metadata={"source": filename, "location_type": "Page", "location_number": page_idx}
                        ))

            # --- 3. Word Document (.docx) ---
            elif file_extension == ".docx":
                doc = DocxDocument(tmp_path)
                full_text = [para.text.strip() for para in doc.paragraphs if para.text.strip()]
                content = "\n".join(full_text)
                if content.strip():
                    all_docs.append(Document(
                        page_content=content,
                        metadata={"source": filename, "location_type": "Section", "location_number": 1}
                    ))

            # --- 4. HTML Web Page (.html, .htm) ---
            elif file_extension in [".html", ".htm"]:
                with open(tmp_path, "r", encoding="utf-8", errors="ignore") as f:
                    soup = BeautifulSoup(f.read(), "html.parser")
                    # Remove scripts and style tags
                    for script in soup(["script", "style"]):
                        script.decompose()
                    text = soup.get_text(separator="\n")
                    lines = [line.strip() for line in text.splitlines() if line.strip()]
                    cleaned_text = "\n".join(lines)
                    if cleaned_text.strip():
                        all_docs.append(Document(
                            page_content=cleaned_text,
                            metadata={"source": filename, "location_type": "Page", "location_number": 1}
                        ))

            # --- 5. Excel (.xlsx, .xls) ---
            elif file_extension in [".xlsx", ".xls"]:
                excel_file = pd.ExcelFile(tmp_path)
                for sheet_idx, sheet_name in enumerate(excel_file.sheet_names, start=1):
                    df = pd.read_excel(excel_file, sheet_name=sheet_name)
                    # Convert table to CSV format text string for clean LLM ingestion
                    content = f"Sheet: {sheet_name}\n" + df.to_csv(index=False)
                    if content.strip():
                        all_docs.append(Document(
                            page_content=content,
                            metadata={"source": filename, "location_type": "Sheet", "location_number": sheet_name}
                        ))

            # --- 6. CSV (.csv) ---
            elif file_extension == ".csv":
                df = pd.read_csv(tmp_path)
                content = df.to_csv(index=False)
                if content.strip():
                    all_docs.append(Document(
                        page_content=content,
                        metadata={"source": filename, "location_type": "Table", "location_number": 1}
                    ))

        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    # Chunk all gathered documents
    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size, 
        chunk_overlap=chunk_overlap
    )
    return text_splitter.split_documents(all_docs)